import os
import json
import time
import tempfile

from dotenv import load_dotenv
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from google import genai
from pypdf import PdfReader
from pptx import Presentation


# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if not GEMINI_API_KEY:
    raise RuntimeError("GEMINI_API_KEY is not set in the .env file")

client = genai.Client(api_key=GEMINI_API_KEY)

AI_MODEL = "gemini-3.5-flash-lite"


# ============================================================
# FASTAPI APP
# ============================================================

app = FastAPI(
    title="LearnSphere AI",
    version="1.0.0",
    description="Personalized Multimodal Learning Companion",
)


# ============================================================
# CORS
# ============================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:3001",
        "http://127.0.0.1:3001",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# BASIC QUESTION MODEL
# ============================================================

class QuestionRequest(BaseModel):
    question: str
    subject: str = "General"
    learning_mode: str = "Beginner"


# ============================================================
# ROOT
# ============================================================

@app.get("/")
def root():
    return {
        "message": "LearnSphere AI API is running"
    }


# ============================================================
# HEALTH
# ============================================================

@app.get("/health")
def health():
    return {
        "status": "healthy"
    }


# ============================================================
# AI LEARNING ASSISTANT
# ============================================================

@app.post("/api/ask")
def ask_question(request: QuestionRequest):

    response = client.models.generate_content(
        model=AI_MODEL,
        contents=f"""
You are LearnSphere AI, a friendly personalized learning companion.

The student selected:

Subject:
{request.subject}

Learning Mode:
{request.learning_mode}

Answer the student's question according to the selected
subject and learning mode.

For Beginner mode:
- Use simple language.
- Explain concepts step by step.
- Give easy examples.

For Intermediate mode:
- Give more technical detail.
- Assume the student already understands the basics.
- Include useful examples.

For Advanced mode:
- Give technically detailed explanations.
- Discuss deeper concepts and practical considerations.

For Exam Preparation mode:
- Focus on important exam concepts.
- Use clear definitions.
- Highlight key points.
- Include short examples where useful.

Student question:
{request.question}
""",
    )

    return {
        "question": request.question,
        "subject": request.subject,
        "learning_mode": request.learning_mode,
        "answer": response.text,
    }


# ============================================================
# NORMAL QUIZ
# ============================================================

class QuizRequest(BaseModel):
    topic: str
    subject: str = "General"
    learning_mode: str = "Beginner"
    previous_score: int = -1


def calculate_adaptive_difficulty(previous_score: int, learning_mode: str):
    if previous_score < 0:
        return "baseline"

    if previous_score < 40:
        return "easier"

    if previous_score < 80:
        return "same"

    return "harder"


@app.post("/api/quiz")
def generate_quiz(request: QuizRequest):

    adaptive_direction = calculate_adaptive_difficulty(
        request.previous_score,
        request.learning_mode
    )

    if adaptive_direction == "easier":
        difficulty_instruction = """
The student's previous score was below 40%.
Make this quiz easier.
Focus on fundamentals, definitions, recognition,
and simple applications.
"""

    elif adaptive_direction == "harder":
        difficulty_instruction = """
The student's previous score was 80% or higher.
Make this quiz harder.
Use deeper reasoning, application questions,
and more challenging distractors.
"""

    elif adaptive_direction == "same":
        difficulty_instruction = """
The student's previous score was between 40% and 79%.
Keep approximately the same difficulty level.
"""

    else:
        difficulty_instruction = """
This is the student's first quiz.
Use a balanced difficulty appropriate for the selected
learning mode.
"""

    response = client.models.generate_content(
        model=AI_MODEL,
        contents=f"""
You are LearnSphere AI, a personalized adaptive quiz generator.

Create exactly 5 multiple-choice questions.

Topic:
{request.topic}

Subject:
{request.subject}

Learning Mode:
{request.learning_mode}

Previous Score:
{request.previous_score}

{difficulty_instruction}

Return ONLY valid JSON.

Use exactly this format:

{{
  "questions": [
    {{
      "question": "Question text",
      "options": [
        "Option A",
        "Option B",
        "Option C",
        "Option D"
      ],
      "correct_answer": 0,
      "explanation": "Short explanation"
    }}
  ]
}}

Important:
- Create exactly 5 questions.
- Each question must have exactly 4 options.
- correct_answer must be a number from 0 to 3.
- Do not include A, B, C, or D inside the options.
- Return valid JSON only.
- Do not use Markdown.
""",
    )

    try:
        quiz_data = json.loads(response.text)
    except json.JSONDecodeError:
        raise HTTPException(
            status_code=500,
            detail="AI returned an invalid quiz format."
        )

    return {
        "topic": request.topic,
        "subject": request.subject,
        "learning_mode": request.learning_mode,
        "adaptive_direction": adaptive_direction,
        "questions": quiz_data["questions"],
    }


# ============================================================
# MATERIAL QUIZ
# ============================================================

class MaterialQuizRequest(BaseModel):
    material_text: str
    subject: str = "General"
    learning_mode: str = "Beginner"
    previous_score: int = -1


@app.post("/api/material-quiz")
def generate_material_quiz(request: MaterialQuizRequest):

    if not request.material_text.strip():
        raise HTTPException(
            status_code=400,
            detail="Material content cannot be empty."
        )

    adaptive_direction = calculate_adaptive_difficulty(
        request.previous_score,
        request.learning_mode
    )

    if adaptive_direction == "easier":
        difficulty_instruction = """
The previous quiz score was below 40%.
Create easier questions focusing on core concepts.
"""

    elif adaptive_direction == "harder":
        difficulty_instruction = """
The previous quiz score was 80% or higher.
Create more challenging questions requiring deeper
understanding and application of the material.
"""

    elif adaptive_direction == "same":
        difficulty_instruction = """
The previous quiz score was between 40% and 79%.
Keep the difficulty approximately the same.
"""

    else:
        difficulty_instruction = """
This is the first quiz from this material.
Use a balanced difficulty appropriate for the selected
learning mode.
"""

    response = client.models.generate_content(
        model=AI_MODEL,
        contents=f"""
You are LearnSphere AI, a personalized adaptive educational
quiz generator.

The student has uploaded learning material.

Create exactly 5 multiple-choice questions using ONLY the
uploaded learning material as the primary source.

Subject:
{request.subject}

Learning Mode:
{request.learning_mode}

Previous Score:
{request.previous_score}

{difficulty_instruction}

IMPORTANT RULES:

- Questions must be supported by the uploaded material.
- Do not invent unsupported facts.
- Create exactly 5 questions.
- Each question must have exactly 4 options.
- Only one option should be correct.
- correct_answer must be a number from 0 to 3.
- Include a short explanation.
- For Beginner mode, keep questions simple.
- For Intermediate mode, use moderate detail.
- For Advanced mode, make questions challenging.
- For Exam Preparation mode, focus on important concepts.
- Return ONLY valid JSON.
- Do not use Markdown.

Use exactly this format:

{{
  "questions": [
    {{
      "question": "Question text",
      "options": [
        "Option A",
        "Option B",
        "Option C",
        "Option D"
      ],
      "correct_answer": 0,
      "explanation": "Short explanation based on the material"
    }}
  ]
}}

UPLOADED LEARNING MATERIAL:
----------------------------
{request.material_text}
----------------------------
""",
    )

    try:
        quiz_data = json.loads(response.text)
    except json.JSONDecodeError:
        raise HTTPException(
            status_code=500,
            detail="AI returned an invalid quiz format."
        )

    if "questions" not in quiz_data:
        raise HTTPException(
            status_code=500,
            detail="Quiz questions were not returned."
        )

    return {
        "subject": request.subject,
        "learning_mode": request.learning_mode,
        "source": "Uploaded learning material",
        "adaptive_direction": adaptive_direction,
        "questions": quiz_data["questions"],
    }


# ============================================================
# IMAGE ANALYSIS
# ============================================================

@app.post("/api/analyze-image")
async def analyze_image(
    file: UploadFile = File(...),
    subject: str = "General",
    learning_mode: str = "Beginner",
):

    image_bytes = await file.read()

    response = client.models.generate_content(
        model=AI_MODEL,
        contents=[
            {
                "inline_data": {
                    "mime_type": file.content_type,
                    "data": image_bytes,
                }
            },
            f"""
You are LearnSphere AI, a friendly personalized learning companion.

Analyze the uploaded image and help the student understand it.

Subject:
{subject}

Learning Mode:
{learning_mode}

Instructions:

- Describe what is visible in the image.
- Explain the important concept step by step.
- For Beginner mode, use simple language and easy examples.
- For Intermediate mode, provide more technical detail.
- For Advanced mode, explain deeper concepts.
- For Exam Preparation mode, focus on important exam points.
- If the image contains a question, solve it clearly.
- If the image contains code, explain what the code does.
- If the image is unclear, honestly say what cannot be determined.

Give a helpful educational explanation.
""",
        ],
    )

    return {
        "filename": file.filename,
        "subject": subject,
        "learning_mode": learning_mode,
        "analysis": response.text,
    }


# ============================================================
# TRACK D - SOURCE MATERIAL UPLOAD
# ============================================================

@app.post("/api/upload-material")
async def upload_material(file: UploadFile = File(...)):

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file was provided."
        )

    filename = file.filename
    extension = os.path.splitext(filename)[1].lower()

    allowed_extensions = {
        ".pdf",
        ".pptx",
        ".txt",
        ".mp4",
        ".webm",
        ".mov",
        ".m4v"
    }

    if extension not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. Please upload "
                "PDF, PPTX, TXT, MP4, WEBM, MOV, or M4V."
            )
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )

    # ========================================================
    # VIDEO
    # ========================================================

    if extension in {
        ".mp4",
        ".webm",
        ".mov",
        ".m4v"
    }:

        mime_types = {
            ".mp4": "video/mp4",
            ".webm": "video/webm",
            ".mov": "video/quicktime",
            ".m4v": "video/mp4",
        }

        mime_type = mime_types.get(
            extension,
            file.content_type or "video/mp4"
        )

        temp_path = None

        try:

            with tempfile.NamedTemporaryFile(
                delete=False,
                suffix=extension
            ) as temp_file:

                temp_file.write(file_bytes)
                temp_path = temp_file.name

            print("Uploading lecture video to Gemini...")

            uploaded_file = client.files.upload(
                file=temp_path
            )

            print(
                "Gemini video file:",
                uploaded_file.name
            )

            # Wait for Gemini video processing.
            while (
                uploaded_file.state
                and uploaded_file.state.name == "PROCESSING"
            ):

                print("Waiting for video processing...")

                time.sleep(3)

                uploaded_file = client.files.get(
                    name=uploaded_file.name
                )

            if (
                uploaded_file.state
                and uploaded_file.state.name == "FAILED"
            ):
                raise HTTPException(
                    status_code=400,
                    detail="Gemini could not process this video."
                )

            video_prompt = """
You are LearnSphere AI.

Analyze this lecture/study video as educational material.

Create a structured lecture transcript/knowledge summary that
can later be used as a source for answering student questions.

IMPORTANT:

1. Identify the major concepts explained in the lecture.
2. Preserve important definitions, examples, formulas, and facts.
3. Organize the information into clear sections.
4. Where possible, identify approximate timestamps such as:
   [Video 00:00-01:00]
   [Video 01:00-02:00]
5. Do not invent timestamps if they cannot be determined.
6. Make the resulting text useful as a searchable learning source.
7. Do not add unrelated information.

Return the educational content in plain text.
"""

            response = client.models.generate_content(
                model=AI_MODEL,
                contents=[
                    uploaded_file,
                    video_prompt
                ]
            )

            extracted_text = response.text.strip()

            if not extracted_text:
                raise HTTPException(
                    status_code=400,
                    detail="No educational content could be extracted from the video."
                )

            return {
                "filename": filename,
                "file_type": "VIDEO",
                "characters": len(extracted_text),
                "source": filename,
                "text": (
                    f"[Lecture Video: {filename}]\n\n"
                    f"{extracted_text}"
                ),
                "message": (
                    "Lecture video processed successfully "
                    "using Gemini multimodal video understanding."
                ),
                "video_supported": True
            }

        except HTTPException:
            raise

        except Exception as error:

            print("VIDEO ERROR:", repr(error))

            raise HTTPException(
                status_code=500,
                detail=f"Could not process lecture video: {str(error)}"
            )

        finally:

            if temp_path and os.path.exists(temp_path):

                try:
                    os.remove(temp_path)
                except Exception:
                    pass

    # ========================================================
    # TEXT / PDF / PPTX
    # ========================================================

    extracted_text = ""

    try:

        # ----------------------------------------------------
        # PDF
        # ----------------------------------------------------

        if extension == ".pdf":

            from io import BytesIO

            pdf_file = BytesIO(file_bytes)

            reader = PdfReader(pdf_file)

            pages = []

            for page_number, page in enumerate(
                reader.pages,
                start=1
            ):

                page_text = page.extract_text() or ""

                if page_text.strip():

                    pages.append(
                        f"[Page {page_number}]\n"
                        f"{page_text.strip()}"
                    )

            extracted_text = "\n\n".join(pages)

        # ----------------------------------------------------
        # PPTX
        # ----------------------------------------------------

        elif extension == ".pptx":

            from io import BytesIO

            presentation = Presentation(
                BytesIO(file_bytes)
            )

            slides = []

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1
            ):

                slide_text = []

                for shape in slide.shapes:

                    if hasattr(shape, "text"):

                        text = shape.text.strip()

                        if text:
                            slide_text.append(text)

                if slide_text:

                    slides.append(
                        f"[Slide {slide_number}]\n"
                        + "\n".join(slide_text)
                    )

            extracted_text = "\n\n".join(slides)

        # ----------------------------------------------------
        # TXT
        # ----------------------------------------------------

        elif extension == ".txt":

            extracted_text = file_bytes.decode(
                "utf-8",
                errors="replace"
            )

    except Exception as error:

        raise HTTPException(
            status_code=400,
            detail=f"Could not read the uploaded file: {str(error)}"
        )

    if not extracted_text.strip():

        raise HTTPException(
            status_code=400,
            detail="No readable text was found in the uploaded file."
        )

    return {
        "filename": filename,
        "file_type": extension.replace(".", "").upper(),
        "characters": len(extracted_text),
        "source": filename,
        "text": extracted_text,
        "message": "Source material extracted successfully.",
        "video_supported": False
    }


# ============================================================
# TRACK D - MATERIAL-BASED AI QUESTION + SOURCE CITATIONS
# ============================================================

class MaterialQuestionRequest(BaseModel):
    question: str
    material_text: str
    subject: str = "General"
    learning_mode: str = "Beginner"


def extract_json_from_response(text: str):

    cleaned = text.strip()

    if cleaned.startswith("```"):
        cleaned = cleaned.replace("```json", "")
        cleaned = cleaned.replace("```", "")
        cleaned = cleaned.strip()

    return json.loads(cleaned)


@app.post("/api/material-ask")
def ask_from_material(request: MaterialQuestionRequest):

    if not request.question.strip():

        raise HTTPException(
            status_code=400,
            detail="Question cannot be empty."
        )

    if not request.material_text.strip():

        raise HTTPException(
            status_code=400,
            detail="Material content cannot be empty."
        )

    response = client.models.generate_content(
        model=AI_MODEL,
        contents=f"""
You are LearnSphere AI, a source-grounded personalized
learning companion.

The student has uploaded learning material.

Answer the student's question using the uploaded material
as the PRIMARY source.

Subject:
{request.subject}

Learning Mode:
{request.learning_mode}

IMPORTANT SOURCE RULES:

- Use only information supported by the uploaded material.
- Do not invent unsupported facts.
- If the answer is not available, say:
  "This information is not available in the uploaded material."
- Identify the exact source location whenever possible.
- The material uses source markers such as:
  [Page 1]
  [Slide 2]
  [Lecture Video: filename]
  [Video 00:10-01:20]
- Cite the relevant source location in the "sources" array.
- Include a short quote or evidence snippet from the material.
- If more than one source location supports the answer,
  include multiple sources.
- For Beginner mode, explain simply.
- For Intermediate mode, give more technical detail.
- For Advanced mode, give deeper technical detail.
- For Exam Preparation mode, focus on key exam concepts.

Return ONLY valid JSON.

Use exactly this structure:

{{
  "answer": "Clear educational answer",
  "sources": [
    {{
      "location": "Slide 2",
      "evidence": "Short supporting text from the material"
    }}
  ]
}}

UPLOADED LEARNING MATERIAL:
----------------------------
{request.material_text}
----------------------------

STUDENT QUESTION:
{request.question}
""",
    )

    try:

        result = extract_json_from_response(
            response.text
        )

    except Exception:

        return {
            "question": request.question,
            "subject": request.subject,
            "learning_mode": request.learning_mode,
            "answer": response.text,
            "sources": [
                {
                    "location": "Uploaded learning material",
                    "evidence": "Source-grounded answer from uploaded material."
                }
            ],
            "source": "Uploaded learning material"
        }

    return {
        "question": request.question,
        "subject": request.subject,
        "learning_mode": request.learning_mode,
        "answer": result.get("answer", ""),
        "sources": result.get("sources", []),
        "source": "Uploaded learning material",
    }